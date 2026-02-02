import os
from docx import Document

DOC_PATH = r"e:\CODE\paizhu-software\backend\muban\派驻检察工作日志.docx"

def restore_inner():
    doc = Document(DOC_PATH)
    table = doc.tables[0]
    
    # Row 2 (Strict / Scene)
    # Col 5 was overwritten by tag {严管}
    try:
        table.cell(2, 5).text = "新增人员\n数量"
        print("Restored (2,5) -> 新增人员")
    except: print("Error restoring (2,5)")
    
    # Row 3 (Equip / Gang / Admission)
    # This row was heavily damaged or confusing. 
    # Based on Screenshot and Row 2 alignment:
    # Row 2: [三大] [现场] [现场Loc] [严管] [新增] [DATA]
    # Row 3 should be:
    #        [警戒] [新增] [DATA]  [涉黑] [收押] [DATA] ??
    # But col counts might differ.
    
    # Let's try to set these headers. 
    # If they overwrite data, it's fine, we re-fill tags later.
    try:
        table.cell(3, 2).text = "警戒具\n检察"
        table.cell(3, 3).text = "新增人员\n数量"
        
        # Shifted to leave space for Tag at Col 4
        table.cell(3, 5).text = "涉黑\n罪犯"
        # Tag at Col 6
        table.cell(3, 7).text = "收押/调\n出数量"
        # Tag at Col 8? (Might error if not exist)
        
        print("Restored Row 3 headers (Shifted).")
    except Exception as e:
        print(f"Error restoring Row 3: {e}")

    doc.save(DOC_PATH)

if __name__ == "__main__":
    restore_inner()
